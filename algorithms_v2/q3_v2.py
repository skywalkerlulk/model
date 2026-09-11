# -*- coding: utf-8 -*-
"""
q3_v2.py — 问题3 判据扫描与输出 (2026 CUMCM A题 v2 方案)
============================================================
输入: results/q2_v2_full.npz (问题2 全流程 60s 网格解, N=96)
判据: max_r C(r,t) < 0.15 kg/kg → 由水分径向单调性, 最严点为 r=0 (中心)
输出: 表5 (每6h × 每0.5cm 至烘干结束), result3.xlsx (60s × 0.1cm)
交叉验证: v1 FVM 参考 t_end = 207960 s (57.77 h)
"""

import sys, os
_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, _ROOT)
os.chdir(_ROOT)

import numpy as np
from openpyxl import Workbook

THRESHOLD = 0.15
TABLE_R_CM = [0.0, 0.5, 1.0, 1.5, 2.0]


def main():
    d = np.load('results/q2_v2_full.npz')
    t, C = d['t'], d['C']                     # (4320,) (4320, 21)
    r_out = np.arange(0.0, 2.01, 0.1)

    # 判据扫描 (最严点为中心: 水分沿半径单调递减 → 中心列即判据序列)
    Cmax = C[:, 0]
    mono = np.all(np.diff(Cmax) <= 1e-12)
    print(f'中心水分单调递减: {mono}')
    meet = np.where(Cmax < THRESHOLD)[0]
    assert len(meet) > 0, '72h 内未达标, 请检查 Q2 求解'
    i_end = meet[0]
    t_end = float(t[i_end])
    print(f'烘干结束判定: t_end = {t_end:.0f} s = {t_end/3600:.4f} h')
    print(f'  结束时刻: 中心 C = {C[i_end, 0]:.6f}, '
          f'上一时刻 (t={t[i_end-1]:.0f}s) = {C[i_end-1, 0]:.6f}')
    print(f'  v1 FVM 参考: 57.77 h; 附件2 收缩平台: 70 h')

    # 表5
    idx = {r: int(round(r / 0.1)) for r in TABLE_R_CM}
    row_times = list(range(6 * 3600, int(t_end), 6 * 3600)) + [int(t_end)]
    print('\n表5  药材烘干过程的水分浓度 (kg/kg)')
    for tt in row_times:
        i = int(tt / 60.0) - 1
        print(f'{tt/3600:.2f}h\t' + '\t'.join(f'{C[i][idx[r]]:.4f}' for r in TABLE_R_CM))

    # result3.xlsx
    wb = Workbook()
    ws = wb.active
    ws.title = 'Sheet1'
    ws.cell(1, 1, '时间\\到药材中心的距离')
    for j, rr in enumerate(r_out):
        ws.cell(1, j + 2, round(rr, 1)).number_format = '0.0'
    for k in range(i_end + 1):
        ws.cell(k + 2, 1, int(t[k]))
        for j, c in enumerate(C[k]):
            cell = ws.cell(k + 2, j + 2, round(float(c), 4))
            cell.number_format = '0.0000'
    wb.save('results/result3.xlsx')
    print(f'\nresult3.xlsx 已保存 ({i_end+1}行 × 21列, 60s 间隔至烘干结束)')


if __name__ == '__main__':
    main()
