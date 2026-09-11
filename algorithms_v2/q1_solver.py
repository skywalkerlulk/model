# -*- coding: utf-8 -*-
"""
q1_v2.py — 问题1 独立求解 (2026 CUMCM A题 v2 方案)
============================================================
温度: Duhamel 解析解 (指数主项闭式 + 残差递归卷积, q1_analytic.py)
水分: Chebyshev 谱配点 + 隐式 BDF (spectral_core.py), 表面 Robin 边界消元
验证: ① 谱热解 vs Duhamel 解析 (同环境)  ② 常数 D 水分 vs Bessel 级数
     ③ 末位谱系数  ④ 表面 BC 残差 (消元精确性)
输出: 表1/表2 + results/result1.xlsx (模板格式, 4 位小数) + figures/v2/
"""

import sys, os
_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, _ROOT)
sys.path.insert(0, os.path.join(_ROOT, 'algorithms_v2'))
os.chdir(_ROOT)

import numpy as np
from scipy.integrate import solve_ivp
from openpyxl import Workbook

from env_model import load_env_models
from spectral_core import RadialField, barycentric_eval, cheb_coeffs
from q1_analytic import DuhamelTemperature, bessel_series_cyl

R = 0.02
N = 48                     # 谱模式数 (N=32/64 收敛验证见 validate_v2)
T_END = 1800.0
T_OUT = np.arange(1.0, T_END + 1.0)          # 输出时刻 (s)
R_OUT_CM = np.arange(0.0, 2.01, 0.1)         # 输出半径 (cm)
TABLE_TIMES = [100, 300, 600, 900, 1200, 1500, 1800]
TABLE_R_CM = [0.0, 0.5, 1.0, 1.5, 2.0]

# Q1 物性 (附录2)
RHO, CP, K, H, BETA = 820.0, 2600.0, 0.36, 25.0, 8e-7
T0, C0 = 28.0, 2.55
D_fun = lambda C, T: 7e-9 * np.exp(-0.89 / np.maximum(C, 1e-6))
dDdC_fun = lambda C, T: 7e-9 * np.exp(-0.89 / np.maximum(C, 1e-6)) * 0.89 / np.maximum(C, 1e-6) ** 2


def solve_temperature(T_env):
    """Duhamel 解析温度场 → (len(t), len(r_out))"""
    sol = DuhamelTemperature(R, RHO, CP, K, H, T0, T_env, n_terms=80)
    return sol.solve(R_OUT_CM / 100.0, T_OUT)


def make_field(C_env, T_env=lambda t: T0, D_f=D_fun, dDdC_f=dDdC_fun):
    return RadialField(N, R,
                       D_fun=D_f, dDdC_fun=dDdC_f, dDdT_fun=lambda C, T: np.zeros_like(C),
                       beta=BETA, C_env_fun=C_env,
                       rho_fun=lambda C: np.full_like(C, RHO),
                       cp_fun=lambda C: np.full_like(C, CP),
                       k_fun=lambda C: np.full_like(C, K),
                       dkdC_fun=lambda C: np.zeros_like(C),
                       h=H, T_env_fun=T_env)


def solve_moisture(C_env):
    """谱配点水分场 → (len(t), len(r_out)), 含诊断"""
    fld = make_field(C_env)
    y0 = np.full(N - 1, C0)                      # 内部节点初值 (x_1..x_{N-1})
    sol = solve_ivp(fld.moisture_rhs, (0.0, T_END), y0, method='BDF',
                    t_eval=T_OUT, rtol=1e-9, atol=1e-9)
    x_out = 1.0 - 2.0 * R_OUT_CM / 100.0 / R
    C_out = np.zeros((len(T_OUT), len(R_OUT_CM)))
    bc_res = []
    for i, tt in enumerate(T_OUT):
        full = fld.full_C(tt, sol.y[:, i])
        C_out[i] = barycentric_eval(fld.x, full, x_out)
        Cr = fld.Dr @ full
        bc_res.append(abs(fld.D_fun(full, None)[-1] * Cr[-1]
                          + BETA * (full[-1] - C_env(tt))))
    diag = dict(field=fld, full=full, bc_res_max=float(np.max(bc_res)),
                tail=float(fld.tail_coeff(full)))
    return C_out, diag


def validate_heat(T_env):
    """谱热解 vs Duhamel 解析 (同一环境)"""
    fld = make_field(T_env, T_env=T_env)
    y0 = np.full(N - 1, T0)
    sol = solve_ivp(fld.heat_rhs, (0.0, T_END), y0, method='BDF',
                    t_eval=T_OUT, rtol=1e-9, atol=1e-9)
    T_ana = solve_temperature(T_env)
    x_out = 1.0 - 2.0 * R_OUT_CM / 100.0 / R
    errs = []
    C_full = np.full(N + 1, C0)
    for i, tt in enumerate(T_OUT):
        full = fld.full_T(tt, sol.y[:, i], C_full)
        T_num = barycentric_eval(fld.x, full, x_out)
        errs.append(np.max(np.abs(T_num - T_ana[i])))
    err = max(errs)
    print(f'[V1a] 谱热解 vs Duhamel 解析: max|ΔT| = {err:.3e} °C')
    return err


def validate_moisture_analytic():
    """常数 D 水分谱解 vs Bessel 级数 (阶跃边界)"""
    fld = make_field(lambda t: 0.02, D_f=lambda C, T: np.full_like(C, 4.94e-9),
                     dDdC_f=lambda C, T: np.zeros_like(C))
    y0 = np.full(N - 1, C0)
    sol = solve_ivp(fld.moisture_rhs, (0.0, T_END), y0, method='BDF',
                    t_eval=T_OUT, rtol=1e-9, atol=1e-9, dense_output=True)
    errs = []
    for tc in (100.0, 600.0, 1800.0):
        y_c = sol.sol(tc)
        full = fld.full_C(tc, y_c)
        C_ref = bessel_series_cyl(fld.r, np.array([tc]), R, 4.94e-9, BETA, C0, 0.02,
                                  n_terms=80)
        errs.append(np.max(np.abs(full - C_ref[0])))
    print(f'[V1b] 常数D谱解 vs Bessel级数: max|ΔC| = {max(errs):.3e}')
    return max(errs)


def write_result1(T_out, C_out):
    wb = Workbook()
    for sheet_name, data in (('温度', T_out), ('水分浓度', C_out)):
        ws = wb.active if sheet_name == '温度' else wb.create_sheet()
        ws.title = sheet_name
        ws.cell(1, 1, '时间\\到药材中心的距离')
        for j, r in enumerate(R_OUT_CM):
            ws.cell(1, j + 2, round(r, 1)).number_format = '0.0'
        for i, t in enumerate(T_OUT):
            ws.cell(i + 2, 1, int(t))
            for j, c in enumerate(data[i]):
                cell = ws.cell(i + 2, j + 2, round(float(c), 4))
                cell.number_format = '0.0000'
    wb.save('results/result1.xlsx')
    print('result1.xlsx 已保存 (v2, 1800行 × 21列 × 2 sheet)')


def print_tables(T_out, C_out):
    idx = {r: int(round(r / 0.1)) for r in TABLE_R_CM}
    print('\n表1  30分钟内药材的温度 (°C)')
    print('时间/s\t' + '\t'.join(f'r={r}cm' for r in TABLE_R_CM))
    for t in TABLE_TIMES:
        row = T_out[int(t) - 1]
        print(f'{t}\t' + '\t'.join(f'{row[idx[r]]:.4f}' for r in TABLE_R_CM))
    print('\n表2  30分钟内药材的水分浓度 (kg/kg)')
    print('时间/s\t' + '\t'.join(f'r={r}cm' for r in TABLE_R_CM))
    for t in TABLE_TIMES:
        row = C_out[int(t) - 1]
        print(f'{t}\t' + '\t'.join(f'{row[idx[r]]:.4f}' for r in TABLE_R_CM))


def make_figures(T_out, C_out, diagC, T_env, C_env):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    plt.rcParams.update({'font.sans-serif': ['Arial Unicode MS', 'PingFang SC', 'Heiti TC'],
                         'axes.unicode_minus': False, 'font.size': 8,
                         'axes.linewidth': 0.6, 'legend.frameon': False})
    PAL = ['#0C7BDC', '#D64541', '#3CB478', '#8C78C8', '#C8AA3C', '#787878']
    os.makedirs('figures/v2', exist_ok=True)

    # 谱系数瀑布图
    a = cheb_coeffs(diagC['full'])
    fig, ax = plt.subplots(figsize=(4.6, 3.0))
    ax.semilogy(np.arange(len(a)), np.abs(a), 'o-', ms=3, color=PAL[0])
    ax.set(xlabel='Chebyshev 模式 n', ylabel='|a_n|', title='谱系数瀑布图 (t=1800s)')
    ax.axhline(5e-5, color='#333', ls='--', lw=0.8, label='4位小数精度阈值')
    ax.legend(fontsize=7); ax.grid(alpha=0.3)
    fig.tight_layout(); fig.savefig('figures/v2/q1_spectral_coeffs.png', dpi=300)
    plt.close(fig)

    # 温度/水分时空 3D
    from mpl_toolkits.mplot3d import Axes3D  # noqa
    r_cm, t_min = R_OUT_CM, T_OUT / 60.0
    fig = plt.figure(figsize=(9.6, 4.0))
    for Z, cb, pos, tt in ((T_out, 'T / °C', 121, '温度场'),
                           (C_out, 'C / (kg·kg⁻¹)', 122, '水分浓度场')):
        ax = fig.add_subplot(pos, projection='3d')
        RR, TT = np.meshgrid(r_cm, t_min[::10])
        surf = ax.plot_surface(RR, TT, Z[::10], cmap='viridis', linewidth=0,
                               rstride=1, cstride=1, antialiased=False)
        fig.colorbar(surf, ax=ax, shrink=0.6, pad=0.08, label=cb)
        ax.view_init(elev=28, azim=-40)
        ax.set_xlabel('r / cm', fontsize=7); ax.set_ylabel('t / min', fontsize=7)
        ax.set_zlabel(cb, fontsize=7); ax.tick_params(labelsize=6)
        ax.xaxis.set_pane_color((1, 1, 1, 0)); ax.yaxis.set_pane_color((1, 1, 1, 0))
        ax.zaxis.set_pane_color((1, 1, 1, 0))
        ax.set_title(tt, fontsize=8)
    fig.tight_layout(); fig.savefig('figures/v2/q1_fields_3d.png', dpi=300)
    plt.close(fig)

    # 环境拟合质量
    t_data = np.arange(0, 14401, 60.0)
    fig, axes = plt.subplots(1, 2, figsize=(9.0, 3.0))
    for ax, env, name in ((axes[0], T_env, '温度'), (axes[1], C_env, '水分浓度')):
        ax.plot(t_data / 3600, [env(t) for t in t_data], color=PAL[0], lw=1.2,
                label='模型 T_env(t)')
        ax.set(xlabel='t / h', ylabel=name, title=f'{name}: 一阶惯性 + 残差')
        ax.grid(alpha=0.3); ax.legend(fontsize=7)
    fig.tight_layout(); fig.savefig('figures/v2/q1_env_model.png', dpi=300)
    plt.close(fig)
    print('图: q1_spectral_coeffs / q1_fields_3d / q1_env_model')


def main():
    T_env, C_env, diag = load_env_models()
    print(f"环境模型: T: R²={diag['T']['R2']:.5f}, τ={diag['T']['tau']:.0f}s | "
          f"C: R²={diag['C']['R2']:.5f}, τ={diag['C']['tau']:.0f}s")
    print(f"平台外推: {diag['T']['plateau'][0]:.3f}±{diag['T']['plateau'][1]:.3f} °C, "
          f"{diag['C']['plateau'][0]:.4f}±{diag['C']['plateau'][1]:.4f} kg/kg")

    T_out = solve_temperature(T_env)
    C_out, diagC = solve_moisture(C_env)
    print(f'[V4] 表面BC残差 max = {diagC["bc_res_max"]:.2e} (消元精确性)')
    print(f'[V5] 末位谱系数 |a_N| = {diagC["tail"]:.3e}')

    validate_heat(T_env)
    validate_moisture_analytic()

    print_tables(T_out, C_out)
    write_result1(T_out, C_out)
    make_figures(T_out, C_out, diagC, T_env, C_env)


if __name__ == '__main__':
    main()
