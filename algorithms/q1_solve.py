# -*- coding: utf-8 -*-
"""
q1_solve.py — 问题1: 预热平衡阶段求解 (表1/表2 + result1.xlsx + 图)
============================================================
模型: 常物性轴对称瞬态传热 (附录2: ρ=820, cp=2600, k=0.36, h=25)
      水分扩散 D(C) = 7e-9·exp(-0.89/C), 对流传质 β=8e-7
环境: 附件1 PCHIP 插值 (0-14400 s)
输出: 表1/表2 (100..1800s × 0..2cm), result1.xlsx (1s × 0.1cm, 4位小数)
"""
import sys, os
_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, _ROOT)      # 兄弟模块导入 (solver_core 等)
os.chdir(_ROOT)                # 数据/结果路径以项目根为基准


import numpy as np
import pandas as pd
from scipy.interpolate import PchipInterpolator
from openpyxl import Workbook
from solver_core import CylinderDryer
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

plt.rcParams['font.sans-serif'] = ['Arial Unicode MS', 'PingFang SC', 'Heiti TC', 'STHeiti']
plt.rcParams['axes.unicode_minus'] = False

R = 0.02            # m
N = 400             # Δr = 0.005 cm → 输出点与节点对齐
DT = 0.25           # s
T_END = 1800.0      # s (表1/表2 窗口)

TABLE_TIMES = [100, 300, 600, 900, 1200, 1500, 1800]
TABLE_R_CM = [0.0, 0.5, 1.0, 1.5, 2.0]
OUT_R_CM = np.arange(0.0, 2.01, 0.1)          # result1.xlsx 列 (21 列)


def load_env():
    df = pd.read_excel('A题/附件/附件1.xlsx')
    t = df['时间'].values.astype(float)
    return (PchipInterpolator(t, df['温度'].values),
            PchipInterpolator(t, df['水分浓度'].values))


def build_solver(T_env, C_env):
    return CylinderDryer(
        R, N,
        rho_fun=lambda C: np.full_like(C, 820.0),
        cp_fun=lambda C: np.full_like(C, 2600.0),
        k_fun=lambda C: np.full_like(C, 0.36),
        D_fun=lambda C, T: 7e-9 * np.exp(-0.89 / np.maximum(C, 1e-6)),
        dDdC_fun=lambda C, T: 7e-9 * np.exp(-0.89 / np.maximum(C, 1e-6)) * 0.89 / np.maximum(C, 1e-6) ** 2,
        dDdT_fun=lambda C, T: np.zeros_like(C),
        h=25.0, beta=8e-7,
        T_env_fun=T_env, C_env_fun=C_env,
        T0=28.0, C0=2.55)


def main():
    T_env, C_env = load_env()
    s = build_solver(T_env, C_env)
    t_out = np.arange(1, int(T_END) + 1, 1)      # 每秒一个快照 (t=1..1800)
    snaps = s.run(T_END, DT, t_out=t_out)
    T_all = np.array([sn[1] for sn in snaps])    # (1800, N+1)
    C_all = np.array([sn[2] for sn in snaps])
    t_axis = np.array([sn[0] for sn in snaps])

    idx = (np.array(TABLE_R_CM) / 100.0 / (R / N)).astype(int)   # 表1/表2 节点下标
    i_t = {t: int(t / 1.0) - 1 for t in TABLE_TIMES}

    # ---------------- 表1 / 表2 ----------------
    print('=' * 60)
    print('表1  30分钟内药材的温度 (°C)')
    print('到药材中心的距离/cm\t' + '\t'.join(f'{r:.1f}' for r in TABLE_R_CM))
    for t in TABLE_TIMES:
        row = T_all[i_t[t]][idx]
        print(f'{t}\t\t\t' + '\t'.join(f'{v:.4f}' for v in row))
    print('-' * 60)
    print('表2  30分钟内药材的水分浓度 (kg/kg)')
    print('到药材中心的距离/cm\t' + '\t'.join(f'{r:.1f}' for r in TABLE_R_CM))
    for t in TABLE_TIMES:
        row = C_all[i_t[t]][idx]
        print(f'{t}\t\t\t' + '\t'.join(f'{v:.4f}' for v in row))

    # ---------------- result1.xlsx ----------------
    out_idx = (OUT_R_CM / 100.0 / (R / N)).astype(int)
    wb = Workbook()
    for sheet_name, data in (('温度', T_all), ('水分浓度', C_all)):
        ws = wb.active if sheet_name == '温度' else wb.create_sheet()
        ws.title = sheet_name
        ws.cell(1, 1, '时间\\到药材中心的距离')
        for j, r in enumerate(OUT_R_CM):
            ws.cell(1, j + 2, round(r, 1)).number_format = '0.0'
        for i, t in enumerate(t_axis):
            ws.cell(i + 2, 1, int(t))
            for j, c in enumerate(data[i][out_idx]):
                cell = ws.cell(i + 2, j + 2, round(float(c), 4))
                cell.number_format = '0.0000'
    wb.save('results/result1.xlsx')
    print('\nresult1.xlsx 已保存 (温度/水分浓度, 1800行 × 21列, 4位小数)')

    # ---------------- 图 ----------------
    r_cm = s.mesh.r * 100.0
    fig, axes = plt.subplots(1, 2, figsize=(11, 4))
    for t in TABLE_TIMES:
        axes[0].plot(r_cm, T_all[i_t[t]], label=f'{t} s')
        axes[1].plot(r_cm, C_all[i_t[t]], label=f'{t} s')
    axes[0].set(xlabel='到药材中心距离 r / cm', ylabel='温度 T / °C', title='预热阶段温度分布')
    axes[1].set(xlabel='到药材中心距离 r / cm', ylabel='水分浓度 C / (kg·kg⁻¹)', title='预热阶段水分浓度分布')
    for ax in axes:
        ax.grid(alpha=.3); ax.legend(fontsize=7, ncol=2)
    fig.tight_layout(); fig.savefig('figures/q1_profiles.png', dpi=200)

    fig, axes = plt.subplots(1, 2, figsize=(11, 4))
    im0 = axes[0].contourf(t_axis / 60.0, r_cm, T_all.T, levels=30, cmap='RdYlBu_r')
    im1 = axes[1].contourf(t_axis / 60.0, r_cm, C_all.T, levels=30, cmap='YlGnBu')
    fig.colorbar(im0, ax=axes[0], label='T / °C')
    fig.colorbar(im1, ax=axes[1], label='C / (kg·kg⁻¹)')
    axes[0].set(xlabel='时间 t / min', ylabel='r / cm', title='温度时空演化')
    axes[1].set(xlabel='时间 t / min', ylabel='r / cm', title='水分浓度时空演化')
    fig.tight_layout(); fig.savefig('figures/q1_heatmaps.png', dpi=200)

    # 环境条件图 (0-14400s)
    df = pd.read_excel('A题/附件/附件1.xlsx')
    fig, axes = plt.subplots(1, 2, figsize=(11, 3.5))
    axes[0].plot(df['时间'] / 3600, df['温度'], '.-', ms=3)
    axes[0].axvline(T_END / 3600, color='r', ls='--', lw=1, label='问题1窗口')
    axes[0].set(xlabel='t / h', ylabel='烘房温度 / °C', title='附件1: 烘房温度'); axes[0].legend(); axes[0].grid(alpha=.3)
    axes[1].plot(df['时间'] / 3600, df['水分浓度'], '.-', ms=3)
    axes[1].axvline(T_END / 3600, color='r', ls='--', lw=1, label='问题1窗口')
    axes[1].set(xlabel='t / h', ylabel='烘房水分浓度 / (kg·kg⁻¹)', title='附件1: 烘房水分浓度'); axes[1].legend(); axes[1].grid(alpha=.3)
    fig.tight_layout(); fig.savefig('figures/q1_env.png', dpi=200)

    print('图已保存: figures/q1_profiles.png, q1_heatmaps.png, q1_env.png')
    np.savez('results/q1_solution.npz', t=t_axis, r=s.mesh.r, T=T_all, C=C_all)
    print('results/q1_solution.npz 已保存')


if __name__ == '__main__':
    main()
