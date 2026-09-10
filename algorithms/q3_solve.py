# -*- coding: utf-8 -*-
"""
q3_solve.py — 问题3: 烘干时长判定 (max C < 0.15 kg/kg)
============================================================
输入: results/q2_full.npz (问题2 全流程 60s 网格解)
判据: max_r C(r,t) < 0.15 → 由水分单调性, 最严点为 r=0 (中心)
输出: 表5 (每6h × 每0.5cm 至烘干结束), result3.xlsx (60s × 0.1cm)
交叉验证: 附件2 半径平台 252000s (70h) 与 t_end 同量级
"""
import sys, os
_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, _ROOT)      # 兄弟模块导入 (solver_core 等)
os.chdir(_ROOT)                # 数据/结果路径以项目根为基准


import numpy as np
from openpyxl import Workbook
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

plt.rcParams['font.sans-serif'] = ['Arial Unicode MS', 'PingFang SC', 'Heiti TC', 'STHeiti']
plt.rcParams['axes.unicode_minus'] = False

R = 0.02
N = 400
THRESHOLD = 0.15
OUT_R_CM = np.arange(0.0, 2.01, 0.1)
TABLE_R_CM = [0.0, 0.5, 1.0, 1.5, 2.0]


def main():
    data = np.load('results/q2_full.npz')
    t, r, C = data['t'], data['r'], data['C']       # (4320,) (401,) (4320,401)
    dr = r[1] - r[0]

    # ---- 判据扫描: max_r C(r,t) < 0.15 ----
    Cmax = C.max(axis=1)
    assert np.all(np.diff(Cmax) <= 1e-12), '中心水分应单调下降, 请检查'
    meet = np.where(Cmax < THRESHOLD)[0]
    assert len(meet) > 0, '72h 内未达标, 需延长模拟'
    i_end = meet[0]
    t_end = float(t[i_end])
    print(f'烘干结束判定: t_end = {t_end} s = {t_end/3600:.4f} h')
    print(f'  结束时中心 C = {C[i_end, 0]:.6f}, 表面 C = {C[i_end, -1]:.6f}')
    print(f'  上一时刻 (t={t[i_end-1]}s) 中心 C = {C[i_end-1, 0]:.6f} (≥0.15, 60s 分辨率内精确)')
    print(f'  附件2 半径平台起点: 252000 s = 70.0 h → 偏差 {(t_end-252000)/3600:.2f} h')

    # ---- 表5: 每6h × 每0.5cm 至结束 ----
    idx = (np.array(TABLE_R_CM) / 100.0 / dr).astype(int)
    row_times = list(range(6 * 3600, int(t_end), 6 * 3600)) + [int(t_end)]
    print('\n表5  药材烘干过程的水分浓度 (kg/kg)')
    print('时间/h\t到药材中心的距离/cm: ' + '\t'.join(f'{r:.1f}' for r in TABLE_R_CM))
    for tt in row_times:
        i = int(tt / 60.0) - 1                    # 60s 网格下标 (t=60,120,...)
        if tt == int(t_end):
            i = i_end
        print(f'{tt/3600:.1f}' + ('' if tt < 3600 else '') + '\t' +
              '\t'.join(f'{v:.4f}' for v in C[i][idx]))

    # ---- result3.xlsx: 60s × 0.1cm, 至烘干结束 ----
    out_idx = (OUT_R_CM / 100.0 / dr).astype(int)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Sheet1'
    ws.cell(1, 1, '时间\\到药材中心的距离')
    for j, rr in enumerate(OUT_R_CM):
        ws.cell(1, j + 2, round(rr, 1)).number_format = '0.0'
    n_rows = i_end + 1                            # t=60,120,...,t_end
    for k in range(n_rows):
        ws.cell(k + 2, 1, int(t[k]))
        for j, c in enumerate(C[k][out_idx]):
            cell = ws.cell(k + 2, j + 2, round(float(c), 4))
            cell.number_format = '0.0000'
    wb.save('results/result3.xlsx')
    print(f'\nresult3.xlsx 已保存 ({n_rows}行 × 21列, 60s间隔至烘干结束)')

    # ---- 图: 水分时间历程 + 判据 ----
    fig, ax = plt.subplots(figsize=(8, 4))
    ax.plot(t / 3600.0, C[:, -1], label='表面 r=2.0 cm', lw=1.2)
    ax.plot(t / 3600.0, C[:, 0], label='中心 r=0', lw=1.2)
    ax.axhline(THRESHOLD, color='r', ls='--', lw=1, label=f'判据 C = {THRESHOLD}')
    ax.axvline(t_end / 3600.0, color='g', ls=':', lw=1.2,
               label=f'烘干结束 t = {t_end/3600:.2f} h')
    ax.set(xlabel='时间 t / h', ylabel='水分浓度 C / (kg·kg⁻¹)',
           title='问题3: 烘干过程水分浓度时间历程')
    ax.legend(); ax.grid(alpha=.3)
    fig.tight_layout(); fig.savefig('figures/q3_drying_time.png', dpi=200)
    print('图已保存: figures/q3_drying_time.png')


if __name__ == '__main__':
    main()
